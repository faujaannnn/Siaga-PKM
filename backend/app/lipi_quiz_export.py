"""Ekspor hasil Kuesioner Kesiapsiagaan Gempa ke Excel (.xlsx)."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.lipi_attempt import ATTEMPT_POST, ATTEMPT_PRE, attempt_label, resolve_user_attempts
from app.routers.lipi_quiz import (
    DEMO_QUESTIONS,
    LIPI_QUESTIONS,
    _LIKERT_4_LABELS,
    _PENGETAHUAN_MAX,
    _QUESTIONNAIRE_VERSION,
    _SIKAP_MAX,
    _SIKAP_POINTS,
    _TD_RAW_MAX,
    _TD_RAW_MAX_POST,
    _likert_4_points,
    _uses_likert_5,
    calculate_lipi_score,
)

_SECTION_LABELS = {
    "dm": "Data Diri",
    "ed": "Efikasi Diri",
    "pg": "Pengetahuan",
    "sk": "Sikap",
    "td": "Tindakan",
}

_SIKAP_LABELS = {
    "ss": "SS (Sangat Setuju)",
    "s": "S (Setuju)",
    "ks": "KS (Kurang Setuju)",
    "ts": "TS (Tidak Setuju)",
    "sts": "STS (Sangat Tidak Setuju)",
}

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_SUBHEADER_FILL = PatternFill("solid", fgColor="D6E4F0")
_SUBHEADER_FONT = Font(bold=True)


def _choice_map(q: dict[str, Any]) -> dict[str, str]:
    if q.get("choices"):
        return {c["id"]: c["text"] for c in q["choices"]}
    if q.get("options"):
        opts = q["options"]
        if q["type"] == "likert_4":
            return {o: _LIKERT_4_LABELS.get(o, o) for o in opts}
        if q["type"] == "likert_5":
            return {o: _SIKAP_LABELS.get(o, o) for o in opts}
        return {o: o for o in opts}
    return {}


def _answer_label(q: dict[str, Any], code: str, attempt_type: str = ATTEMPT_PRE) -> str:
    if not code:
        return ""
    if q["type"] == "text":
        return str(code)
    if _uses_likert_5(q["section"], attempt_type):
        return _SIKAP_LABELS.get(code, code)
    return _choice_map(q).get(code, code)


def _item_points(q: dict[str, Any], code: str, attempt_type: str = ATTEMPT_PRE) -> float | str:
    if not code:
        return ""
    if q["section"] == "pg":
        correct = q.get("correct", "")
        return 1.0 if code == correct else 0.0
    if _uses_likert_5(q["section"], attempt_type):
        return float(_SIKAP_POINTS.get(code, 0))
    if q["section"] in ("ed", "td"):
        return _likert_4_points(code)
    return ""


def _format_dt(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        dt = value
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        local = dt.astimezone(timezone.utc)
        return local.strftime("%Y-%m-%d %H:%M:%S UTC")
    return str(value)


def _style_header_row(ws, row: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize_columns(ws, max_width: int = 48) -> None:
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        width = 0
        for cell in col_cells:
            if cell.value is None:
                continue
            width = max(width, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = max(width, 10)


def _write_sheet_rows(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    _autosize_columns(ws)


def _build_rumus_rows() -> list[list[str]]:
    return [
        ["Bagian", "Variabel", "Rumus / Aturan", "Rentang Skor", "Kategori"],
        [
            "Pengetahuan (pg)",
            "pengetahuan_score",
            f"({{benar}} / {_PENGETAHUAN_MAX}) × 100%",
            "0–100%",
            "Baik ≥76%, Cukup 56–75%, Kurang ≤55%",
        ],
        [
            "Pengetahuan (pg)",
            "poin per butir",
            "1 jika jawaban = kunci benar, 0 jika salah",
            f"0–1 per butir (maks {_PENGETAHUAN_MAX})",
            "—",
        ],
        [
            "Efikasi Diri (ed) · pre-test",
            "poin per butir",
            "poin = 5 − jawaban (Likert 1=sangat setuju … 4=tidak setuju)",
            "1–4 poin per butir",
            "—",
        ],
        [
            "Efikasi Diri (ed) · post-test",
            "poin per butir",
            "SS=5, S=4, KS=3, TS=2, STS=1 (sama seperti sikap)",
            "1–5 poin per butir",
            "—",
        ],
        [
            "Efikasi Diri (ed)",
            "efikasi_diri_raw",
            "jumlah poin 15 butir",
            "Pre: 15–60 · Post: 15–75",
            "—",
        ],
        [
            "Efikasi Diri (ed)",
            "efikasi_diri_score",
            "10 + (raw − 15) × (30 / span); span=45 (pre) atau 60 (post)",
            "10–40 (skala GSES)",
            "Tinggi 30–40, Sedang 20–29, Rendah 10–19",
        ],
        [
            "Sikap (sk)",
            "poin per butir",
            "SS=5, S=4, KS=3, TS=2, STS=1",
            "1–5 poin per butir",
            "—",
        ],
        [
            "Sikap (sk)",
            "sikap_score",
            f"(jumlah poin / {_SIKAP_MAX}) × 100%",
            "0–100%",
            "Baik ≥76%, Cukup 56–75%, Kurang ≤55%",
        ],
        [
            "Tindakan (td) · pre-test",
            "poin per butir",
            "poin = 5 − jawaban (Likert 1–4)",
            "1–4 poin per butir",
            "—",
        ],
        [
            "Tindakan (td) · post-test",
            "poin per butir",
            "SS=5, S=4, KS=3, TS=2, STS=1",
            "1–5 poin per butir",
            "—",
        ],
        [
            "Tindakan (td)",
            "tindakan_score",
            f"(jumlah poin / maks) × 100%; maks pre={_TD_RAW_MAX}, post={_TD_RAW_MAX_POST}",
            "0–100%",
            "Baik ≥50%, Buruk <50%",
        ],
        [
            "Data Diri (dm)",
            "—",
            "Tidak dinilai; hanya untuk profil responden",
            "—",
            "—",
        ],
        [
            "Meta",
            "versi kuesioner",
            _QUESTIONNAIRE_VERSION,
            "—",
            "—",
        ],
        [
            "Alur penelitian",
            "attempt_type",
            "Pengisian 1 = pre_test (ed/td Likert 1–4), pengisian 2 = post_test (semua Likert SS–STS)",
            "pre_test | post_test",
            "Status: belum · post_menunggu · lengkap",
        ],
    ]


def _build_bank_soal_rows() -> list[list[Any]]:
    rows: list[list[Any]] = []
    for q in DEMO_QUESTIONS + LIPI_QUESTIONS:
        choices = _choice_map(q)
        choice_text = "; ".join(f"{k}={v}" for k, v in choices.items()) if choices else ""
        rows.append([
            q["id"],
            _SECTION_LABELS.get(q["section"], q["section"]),
            q.get("number", ""),
            q["text"],
            q["type"],
            q.get("correct", ""),
            choice_text,
        ])
    return rows


def _doc_attempt_type(doc: dict[str, Any], index_in_sorted: int) -> str:
    stored = doc.get("attempt_type")
    if stored in (ATTEMPT_PRE, ATTEMPT_POST):
        return str(stored)
    return ATTEMPT_PRE if index_in_sorted == 0 else ATTEMPT_POST


def _session_scores(doc: dict[str, Any], answers: dict[str, str]) -> dict[str, Any]:
    if answers:
        attempt = str(doc.get("attempt_type") or "pre_test")
        return calculate_lipi_score(answers, attempt)
    return {
        "pengetahuan_score": doc.get("pengetahuan_score", 0),
        "pengetahuan_earned": doc.get("pengetahuan_earned", 0),
        "pengetahuan_max": doc.get("pengetahuan_max", _PENGETAHUAN_MAX),
        "pengetahuan_level": doc.get("pengetahuan_level", doc.get("level", "")),
        "efikasi_diri_raw": doc.get("efikasi_diri_raw", ""),
        "efikasi_diri_score": doc.get("efikasi_diri_score", ""),
        "efikasi_diri_level": doc.get("efikasi_diri_level", ""),
        "sikap_score": doc.get("sikap_score", ""),
        "sikap_raw": doc.get("sikap_raw", ""),
        "sikap_max": doc.get("sikap_max", _SIKAP_MAX),
        "sikap_level": doc.get("sikap_level", ""),
        "tindakan_score": doc.get("tindakan_score", ""),
        "tindakan_raw": doc.get("tindakan_raw", ""),
        "tindakan_max": doc.get("tindakan_max", _TD_RAW_MAX),
        "tindakan_level": doc.get("tindakan_level", ""),
    }


def _build_ringkasan_rows(
    sessions: list[dict[str, Any]],
    user_map: dict[str, dict[str, str]],
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "No",
        "Session ID",
        "Username",
        "Nama Lengkap",
        "Puskesmas",
        "Tipe Attempt",
        "Waktu Submit",
        "Versi Kuesioner",
        "Pengetahuan (%)",
        "Pengetahuan Level",
        "Benar",
        "Maks",
        "Efikasi Diri (GSES)",
        "Efikasi Raw",
        "Efikasi Level",
        "Sikap (%)",
        "Sikap Raw",
        "Sikap Maks",
        "Sikap Level",
        "Tindakan (%)",
        "Tindakan Raw",
        "Tindakan Maks",
        "Tindakan Level",
    ]
    rows: list[list[Any]] = []
    by_user: dict[str, list[dict[str, Any]]] = {}
    for doc in sessions:
        by_user.setdefault(str(doc.get("username") or ""), []).append(doc)
    for username, user_docs in by_user.items():
        user_docs.sort(key=lambda d: str(d.get("submitted_at") or ""))

    flat_sessions: list[tuple[int, dict[str, Any], int]] = []
    for user_docs in by_user.values():
        for i, doc in enumerate(user_docs):
            flat_sessions.append((i, doc, i))

    flat_sessions.sort(key=lambda t: str(t[1].get("submitted_at") or ""))

    for idx, (user_idx, doc, _) in enumerate(flat_sessions, start=1):
        username = str(doc.get("username") or "")
        user = user_map.get(username, {})
        answers = dict(doc.get("answers") or {})
        scores = _session_scores(doc, answers)
        atype = _doc_attempt_type(doc, user_idx)
        rows.append([
            idx,
            str(doc.get("_id", "")),
            username,
            user.get("display_name", ""),
            user.get("puskesmas_nama", ""),
            attempt_label(atype),
            _format_dt(doc.get("submitted_at")),
            doc.get("questionnaire_version", _QUESTIONNAIRE_VERSION),
            scores.get("pengetahuan_score", ""),
            scores.get("pengetahuan_level", ""),
            scores.get("pengetahuan_earned", ""),
            scores.get("pengetahuan_max", ""),
            scores.get("efikasi_diri_score", ""),
            scores.get("efikasi_diri_raw", ""),
            scores.get("efikasi_diri_level", ""),
            scores.get("sikap_score", ""),
            scores.get("sikap_raw", ""),
            scores.get("sikap_max", ""),
            scores.get("sikap_level", ""),
            scores.get("tindakan_score", ""),
            scores.get("tindakan_raw", ""),
            scores.get("tindakan_max", ""),
            scores.get("tindakan_level", ""),
        ])
    return headers, rows


def _build_jawaban_rows(
    sessions: list[dict[str, Any]],
    user_map: dict[str, dict[str, str]],
) -> tuple[list[str], list[list[Any]]]:
    all_questions = DEMO_QUESTIONS + LIPI_QUESTIONS
    base_headers = [
        "No",
        "Session ID",
        "Username",
        "Nama Lengkap",
        "Puskesmas",
        "Tipe Attempt",
        "Waktu Submit",
    ]
    question_headers: list[str] = []
    for q in all_questions:
        qid = q["id"]
        question_headers.append(f"{qid}_kode")
        question_headers.append(f"{qid}_jawaban")
        if q["section"] != "dm":
            question_headers.append(f"{qid}_poin")
            if q["section"] == "pg":
                question_headers.append(f"{qid}_benar")

    headers = base_headers + question_headers
    rows: list[list[Any]] = []

    by_user_j: dict[str, list[dict[str, Any]]] = {}
    for doc in sessions:
        by_user_j.setdefault(str(doc.get("username") or ""), []).append(doc)
    for docs in by_user_j.values():
        docs.sort(key=lambda d: str(d.get("submitted_at") or ""))

    flat_j: list[tuple[int, dict[str, Any]]] = []
    for docs in by_user_j.values():
        for i, doc in enumerate(docs):
            flat_j.append((i, doc))
    flat_j.sort(key=lambda t: str(t[1].get("submitted_at") or ""))

    for idx, (user_idx, doc) in enumerate(flat_j, start=1):
        username = str(doc.get("username") or "")
        user = user_map.get(username, {})
        demographics = dict(doc.get("demographics") or {})
        answers = dict(doc.get("answers") or {})
        atype = _doc_attempt_type(doc, user_idx)

        row: list[Any] = [
            idx,
            str(doc.get("_id", "")),
            username,
            user.get("display_name", ""),
            user.get("puskesmas_nama", ""),
            attempt_label(atype),
            _format_dt(doc.get("submitted_at")),
        ]

        for q in all_questions:
            qid = q["id"]
            if q["section"] == "dm":
                code = demographics.get(qid, "")
            else:
                code = answers.get(qid, "")
            label = _answer_label(q, str(code), atype)
            row.append(code)
            row.append(label)
            if q["section"] != "dm":
                pts = _item_points(q, str(code), atype)
                row.append(pts)
                if q["section"] == "pg":
                    row.append("Ya" if pts == 1.0 else "Tidak" if code else "")

        rows.append(row)

    return headers, rows


def _build_long_format_rows(
    sessions: list[dict[str, Any]],
    user_map: dict[str, dict[str, str]],
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Session ID",
        "Username",
        "Nama Lengkap",
        "Puskesmas",
        "Tipe Attempt",
        "Waktu Submit",
        "Bagian",
        "ID Soal",
        "No",
        "Pertanyaan",
        "Kode Jawaban",
        "Jawaban",
        "Poin",
        "Benar (pg)",
    ]
    rows: list[list[Any]] = []
    all_questions = DEMO_QUESTIONS + LIPI_QUESTIONS

    by_user_l: dict[str, list[dict[str, Any]]] = {}
    for doc in sessions:
        by_user_l.setdefault(str(doc.get("username") or ""), []).append(doc)
    for docs in by_user_l.values():
        docs.sort(key=lambda d: str(d.get("submitted_at") or ""))

    for username, user_docs in by_user_l.items():
        user = user_map.get(username, {})
        for user_idx, doc in enumerate(user_docs):
            demographics = dict(doc.get("demographics") or {})
            answers = dict(doc.get("answers") or {})
            session_id = str(doc.get("_id", ""))
            submitted = _format_dt(doc.get("submitted_at"))
            raw_atype = _doc_attempt_type(doc, user_idx)
            atype = attempt_label(raw_atype)

            for q in all_questions:
                qid = q["id"]
                if q["section"] == "dm":
                    code = demographics.get(qid, "")
                else:
                    code = answers.get(qid, "")
                label = _answer_label(q, str(code), raw_atype)
                pts = _item_points(q, str(code), raw_atype) if q["section"] != "dm" else ""
                benar = ""
                if q["section"] == "pg" and code:
                    benar = "Ya" if pts == 1.0 else "Tidak"
                rows.append([
                    session_id,
                    username,
                    user.get("display_name", ""),
                    user.get("puskesmas_nama", ""),
                    atype,
                    submitted,
                    _SECTION_LABELS.get(q["section"], q["section"]),
                    qid,
                    q.get("number", ""),
                    q["text"],
                    code,
                    label,
                    pts,
                    benar,
                ])

    return headers, rows


def _delta(post: Any, pre: Any) -> str | float:
    if post == "" or pre == "" or post is None or pre is None:
        return ""
    try:
        return round(float(post) - float(pre), 1)
    except (TypeError, ValueError):
        return ""


def _build_pre_post_compare_rows(
    sessions: list[dict[str, Any]],
    user_map: dict[str, dict[str, str]],
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "No",
        "Username",
        "Nama Lengkap",
        "Puskesmas",
        "Status Kuesioner",
        "Pre - Pengetahuan (%)",
        "Post - Pengetahuan (%)",
        "Δ Pengetahuan",
        "Pre - Efikasi (/40)",
        "Post - Efikasi (/40)",
        "Δ Efikasi",
        "Pre - Sikap (%)",
        "Post - Sikap (%)",
        "Δ Sikap",
        "Pre - Tindakan (%)",
        "Post - Tindakan (%)",
        "Δ Tindakan",
        "Waktu Pre-test",
        "Waktu Post-test",
    ]
    by_user: dict[str, list[dict[str, Any]]] = {}
    for doc in sessions:
        by_user.setdefault(str(doc.get("username") or ""), []).append(doc)

    rows: list[list[Any]] = []
    for no, (username, user_docs) in enumerate(sorted(by_user.items()), start=1):
        user = user_map.get(username, {})
        state = resolve_user_attempts(user_docs)
        pre = _session_scores(state["pre_session"] or {}, dict((state["pre_session"] or {}).get("answers") or {})) if state["pre_session"] else {}
        post = _session_scores(state["post_session"] or {}, dict((state["post_session"] or {}).get("answers") or {})) if state["post_session"] else {}
        status_map = {
            "belum": "Belum pre-test",
            "post_menunggu": "Pre-test selesai · post-test menunggu",
            "lengkap": "Pre-test & post-test selesai",
        }
        rows.append([
            no,
            username,
            user.get("display_name", ""),
            user.get("puskesmas_nama", ""),
            status_map.get(str(state["kuis_status"]), str(state["kuis_status"])),
            pre.get("pengetahuan_score", ""),
            post.get("pengetahuan_score", ""),
            _delta(post.get("pengetahuan_score", ""), pre.get("pengetahuan_score", "")),
            pre.get("efikasi_diri_score", ""),
            post.get("efikasi_diri_score", ""),
            _delta(post.get("efikasi_diri_score", ""), pre.get("efikasi_diri_score", "")),
            pre.get("sikap_score", ""),
            post.get("sikap_score", ""),
            _delta(post.get("sikap_score", ""), pre.get("sikap_score", "")),
            pre.get("tindakan_score", ""),
            post.get("tindakan_score", ""),
            _delta(post.get("tindakan_score", ""), pre.get("tindakan_score", "")),
            _format_dt((state["pre_session"] or {}).get("submitted_at")),
            _format_dt((state["post_session"] or {}).get("submitted_at")),
        ])
    return headers, rows


def build_lipi_quiz_xlsx(
    sessions: list[dict[str, Any]],
    user_map: dict[str, dict[str, str]],
) -> bytes:
    """Bangun file Excel multi-sheet dari daftar sesi kuesioner."""
    wb = Workbook()

    ws_rumus = wb.active
    ws_rumus.title = "Rumus Skor"
    rumus_rows = _build_rumus_rows()
    ws_rumus.append(["Kuesioner Kesiapsiagaan Gempa — Rumus Perhitungan Skor"])
    ws_rumus.merge_cells("A1:E1")
    ws_rumus["A1"].font = Font(bold=True, size=13)
    ws_rumus.append([])
    for row in rumus_rows:
        ws_rumus.append(row)
    _style_header_row(ws_rumus, 3, 5)
    _autosize_columns(ws_rumus)

    ws_bank = wb.create_sheet("Bank Soal")
    bank_headers = ["ID", "Bagian", "No", "Pertanyaan", "Tipe", "Kunci Benar", "Pilihan"]
    _write_sheet_rows(ws_bank, bank_headers, _build_bank_soal_rows())

    ws_compare = wb.create_sheet("Pre vs Post")
    compare_headers, compare_rows = _build_pre_post_compare_rows(sessions, user_map)
    _write_sheet_rows(ws_compare, compare_headers, compare_rows)

    ws_ringkasan = wb.create_sheet("Ringkasan Skor")
    ringkasan_headers, ringkasan_rows = _build_ringkasan_rows(sessions, user_map)
    _write_sheet_rows(ws_ringkasan, ringkasan_headers, ringkasan_rows)

    ws_jawaban = wb.create_sheet("Jawaban Lengkap")
    jawaban_headers, jawaban_rows = _build_jawaban_rows(sessions, user_map)
    _write_sheet_rows(ws_jawaban, jawaban_headers, jawaban_rows)

    ws_long = wb.create_sheet("Jawaban per Baris")
    long_headers, long_rows = _build_long_format_rows(sessions, user_map)
    _write_sheet_rows(ws_long, long_headers, long_rows)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

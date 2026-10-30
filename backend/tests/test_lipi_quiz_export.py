"""Unit test ekspor Excel kuesioner LIPI."""

from datetime import datetime, timezone

from openpyxl import load_workbook

from app.lipi_quiz_export import build_lipi_quiz_xlsx


def _sample_session() -> dict:
    return {
        "_id": "sess001",
        "username": "staff01",
        "submitted_at": datetime(2026, 3, 10, 12, 0, tzinfo=timezone.utc),
        "questionnaire_version": "gempa_puskesmas_v2",
        "demographics": {
            "dm1": "laki_laki",
            "dm2": "32",
            "dm3": "s1",
            "dm4": "Perawat",
            "dm5": "5 tahun",
            "dm6": "ya",
            "dm7": "tidak",
            "dm8": "ya",
        },
        "answers": {
            **{f"ed{i}": "1" for i in range(1, 16)},
            "pg1": "b",
            "pg2": "a",
            "pg3": "a",
            "pg4": "b",
            "pg5": "b",
            "pg6": "a",
            "pg7": "c",
            "pg8": "a",
            **{f"sk{i}": "ss" for i in range(1, 16)},
            **{f"td{i}": "2" for i in range(1, 16)},
        },
        "pengetahuan_score": 100.0,
        "pengetahuan_level": "Baik",
        "efikasi_diri_score": 40.0,
        "efikasi_diri_level": "Tinggi",
        "sikap_score": 100.0,
        "sikap_level": "Baik",
        "tindakan_score": 75.0,
        "tindakan_level": "Baik",
    }


def test_build_lipi_quiz_xlsx_has_expected_sheets():
    data = build_lipi_quiz_xlsx(
        [_sample_session()],
        {"staff01": {"display_name": "Budi", "puskesmas_nama": "Puskesmas A"}},
    )
    wb = load_workbook(filename=__import__("io").BytesIO(data))
    assert wb.sheetnames == [
        "Rumus Skor",
        "Bank Soal",
        "Pre vs Post",
        "Ringkasan Skor",
        "Jawaban Lengkap",
        "Jawaban per Baris",
    ]
    ringkasan = wb["Ringkasan Skor"]
    assert ringkasan.max_row >= 2
    assert ringkasan["C2"].value == "staff01"
    assert ringkasan["D2"].value == "Budi"

    rumus = wb["Rumus Skor"]
    assert "Pengetahuan" in str(rumus["A4"].value)

    bank = wb["Bank Soal"]
    assert bank["A2"].value == "dm1"
    assert bank["A2"].value is not None
